"""
Completa el modelo de avión y la matrícula en vuelos_BER_2026-10-22_y_29.xlsx usando
la API pública del sitio del aeropuerto BER (la misma que usa su página de salidas/llegadas).

Uso:
  pip install requests openpyxl
  python completar_aviones_ber.py vuelos_BER_2026-10-22_y_29.xlsx [vuelos_ber_api.json]

Si pasás el JSON descargado desde el navegador (snippet de consola), usa ese archivo
en lugar de consultar la API desde Python.

Genera <archivo>_con_aviones.xlsx con:
  - columna I "Modelo de avión" completa, columna K "Matrícula",
  - columnas L "Asientos (típ.)" y M "Estela ICAO" (fórmulas a la hoja "Aeronaves"),
  - hoja "Aeronaves": modelo -> asientos típicos -> estela ICAO (editable).
Se puede volver a correr sobre el archivo _con_aviones.xlsx: completa lo que falte y
conserva los valores que hayas editado a mano en "Aeronaves".
"""
import re, sys, time, datetime as dt
import requests, openpyxl
from copy import copy
from openpyxl.styles import Alignment

API = "https://ber.berlin-airport.de/api.flights.json"
SHEET, FIRST_ROW = "Vuelos BER", 5
PER_PAGE = 100     # si la API no acepta tanto, igual se recorren todas las páginas
PAUSE = 0.7        # segundos entre consultas, para no sobrecargar el sitio
HEADERS = {"User-Agent": "Mozilla/5.0 (proyecto académico; consultas lentas)",
           "Accept": "application/json"}

# Asientos típicos (configuración habitual del modelo, no la de cada aerolínea) y estela ICAO.
TYPES = [
    (r"a380", 500, "J"), (r"a350-?1000", 350, "H"), (r"a350", 310, "H"), (r"a340", 270, "H"),
    (r"a330.*neo|a330-?9", 290, "H"), (r"a330", 280, "H"), (r"777", 330, "H"),
    (r"787-?10", 330, "H"), (r"787-?9", 290, "H"), (r"787", 250, "H"), (r"767", 220, "H"),
    (r"757", 190, "M"),
    (r"737.*max ?9|737-9", 193, "M"), (r"737.*(8-?200|max ?200)|8200", 197, "M"),
    (r"737.*max|737-8\b|737 8\b|737-8 ", 189, "M"), (r"737-?9", 180, "M"),
    (r"737-?8", 189, "M"), (r"737-?7", 149, "M"), (r"737", 150, "M"),
    (r"a321.*neo|a321n|a21n", 220, "M"), (r"a321", 200, "M"),
    (r"a320.*neo|a320n|a20n", 186, "M"), (r"a320", 180, "M"),
    (r"a319.*neo|a19n", 150, "M"), (r"a319", 144, "M"), (r"a318", 110, "M"),
    (r"a220-?3|bcs3|cs300", 145, "M"), (r"a220|bcs1|cs100", 120, "M"),
    (r"195.*e2|e295", 136, "M"), (r"195", 120, "M"), (r"190.*e2|e290", 100, "M"),
    (r"190", 100, "M"), (r"175", 80, "M"), (r"170", 76, "M"),
    (r"crj.?9|crj.?1000", 90, "M"), (r"crj.?7", 70, "M"), (r"crj", 50, "M"),
    (r"atr.?72", 70, "M"), (r"atr.?42", 48, "M"), (r"dash ?8|dhc-?8|q400", 78, "M"),
    (r"saab ?2000", 50, "M"), (r"saab|340", 34, "M"), (r"328", 32, "M"),
]
def classify(model):
    m = (model or "").lower()
    for pat, seats, wtc in TYPES:
        if re.search(pat, m):
            return seats, wtc
    return None, None

def norm(flight):
    """'EW 8904' / 'EW8904' / 'EW 08904' -> ('EW', '8904')"""
    f = re.sub(r"\s+", "", flight or "").upper()
    mm = re.match(r"^([A-Z0-9]{2}[A-Z]?)(\d+)[A-Z]?$", f)
    if not mm:
        return f, f
    pre, num = mm.group(1), mm.group(2).lstrip("0") or "0"
    if pre[-1].isdigit() and len(pre) == 3:      # p.ej. '6H2' mal partido
        pre, num = pre[:2], (pre[2] + num).lstrip("0")
    return pre, num

def fetch_day(day, direction):
    nxt = day + dt.timedelta(days=1)
    items, page = [], 1
    while True:
        params = dict(arrivalDeparture=direction, dateFrom=f"{day.isoformat()}T00:00:00",
                      dateUntil=nxt.isoformat(), search="", lang="en", page=page,
                      terminal="", itemsPerPage=PER_PAGE)
        r = requests.get(API, params=params, headers=HEADERS, timeout=60)
        r.raise_for_status()
        data = (r.json() or {}).get("data") or {}
        items += data.get("items") or []
        total_pages = data.get("total_pages") or 1
        if page >= total_pages:
            return items
        page += 1
        time.sleep(PAUSE)

def load_json(json_path):
    import json
    data = json.load(open(json_path, encoding="utf-8"))
    out = {}
    for it in data:
        d = dt.date.fromisoformat((it.get("scheduled_time") or "")[:10])
        tipo = "Salida" if it.get("is_departure") else "Llegada"
        out.setdefault((d, tipo), []).append(it)
    return out

def main(path, json_path=None):
    wb = openpyxl.load_workbook(path)
    ws = wb[SHEET]
    last = ws.max_row
    days = sorted({ws.cell(r, 1).value.date() for r in range(FIRST_ROW, last + 1)})

    from_json = load_json(json_path) if json_path else None
    exact, by_num = {}, {}
    for d in days:
        for direction, tipo in (("D", "Salida"), ("A", "Llegada")):
            if from_json is not None:
                its = from_json.get((d, tipo), [])
            else:
                try:
                    its = fetch_day(d, direction)
                except Exception as e:
                    print(f"  {d} {tipo}: error al consultar ({e})"); continue
            con_tipo = 0
            for it in its:
                st = (it.get("scheduled_time") or "")
                if st[:10] != d.isoformat():
                    continue
                model, reg = it.get("aircraft_type") or "", it.get("aircraft_reg") or ""
                if model: con_tipo += 1
                pre, num = norm(it.get("flight_number"))
                rec = (model, reg)
                exact[(d, tipo, pre, num, st[11:16])] = rec
                exact.setdefault((d, tipo, "*", num, st[11:16]), rec)
                by_num.setdefault((d, tipo, num), []).append(rec)
            print(f"  {d} {tipo}: {len(its)} vuelos en la API, {con_tipo} con modelo de avión")
            if from_json is None:
                time.sleep(PAUSE)

    # Encabezados K, L, M (idempotente)
    hdr = ws.cell(4, 10)
    for col, h, w in ((11, "Matrícula", 12), (12, "Asientos (típ.)", 14), (13, "Estela ICAO", 12)):
        c = ws.cell(4, col, h)
        c.font, c.fill = copy(hdr.font), copy(hdr.fill)
        c.alignment, c.border = copy(hdr.alignment), copy(hdr.border)
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = w

    filled = pending = 0
    for r in range(FIRST_ROW, last + 1):
        d = ws.cell(r, 1).value.date(); tipo = ws.cell(r, 2).value
        hhmm = ws.cell(r, 3).value; pre, num = norm(ws.cell(r, 4).value)
        rec = exact.get((d, tipo, pre, num, hhmm)) or exact.get((d, tipo, "*", num, hhmm))
        if not rec:
            cands = {x for x in by_num.get((d, tipo, num), []) if x[0]}
            rec = next(iter(cands)) if len(cands) == 1 else None
        if rec and rec[0]:
            ws.cell(r, 9).value = rec[0]
            if rec[1]: ws.cell(r, 11).value = rec[1]
        if ws.cell(r, 9).value: filled += 1
        else: pending += 1

    # Hoja Aeronaves (conserva ediciones manuales previas)
    prev = {}
    if "Aeronaves" in wb.sheetnames:
        old = wb["Aeronaves"]
        for row in old.iter_rows(min_row=2, values_only=True):
            if row and row[0] and (row[1] is not None or row[2] is not None):
                prev[row[0]] = (row[1], row[2])
        del wb["Aeronaves"]
    wa = wb.create_sheet("Aeronaves")
    for c, h in enumerate(["Modelo de avión", "Asientos (típ.)", "Estela ICAO", "Vuelos", "Nota"], 1):
        cell = wa.cell(1, c, h)
        cell.font, cell.fill = copy(hdr.font), copy(hdr.fill)
        cell.alignment, cell.border = copy(hdr.alignment), copy(hdr.border)
    models = sorted({ws.cell(r, 9).value for r in range(FIRST_ROW, last + 1) if ws.cell(r, 9).value})
    body = ws.cell(FIRST_ROW, 4)
    sin = []
    for i, m in enumerate(models, 2):
        seats, wtc = prev.get(m) or classify(m)
        vals = [m, seats, wtc, f"=COUNTIF('{SHEET}'!$I${FIRST_ROW}:$I${last},A{i})",
                None if seats else "Sin clasificar: completar a mano"]
        if not seats: sin.append(m)
        for c, v in enumerate(vals, 1):
            cell = wa.cell(i, c, v); cell.font = copy(body.font); cell.border = copy(body.border)
    n = max(len(models) + 1, 2)
    wa.cell(n + 2, 1, "Asientos = configuración típica del modelo, no la de cada aerolínea. "
                      "Estela ICAO (L/M/H/J). Editable: la tabla principal toma estos valores con fórmulas."
            ).font = copy(body.font)
    for col, w in zip("ABCDE", (30, 15, 12, 9, 32)):
        wa.column_dimensions[col].width = w

    # Fórmulas L y M + estilo K, L, M
    for r in range(FIRST_ROW, last + 1):
        ws.cell(r, 12).value = (f'=IF($I{r}="","",IFERROR(INDEX(Aeronaves!$B$2:$B${n},'
                                f'MATCH($I{r},Aeronaves!$A$2:$A${n},0)),""))')
        ws.cell(r, 13).value = (f'=IF($I{r}="","",IFERROR(INDEX(Aeronaves!$C$2:$C${n},'
                                f'MATCH($I{r},Aeronaves!$A$2:$A${n},0)),""))')
        s = ws.cell(r, 10)
        for c in (11, 12, 13):
            cell = ws.cell(r, c)
            cell.font, cell.fill, cell.border = copy(s.font), copy(s.fill), copy(s.border)
            cell.alignment = Alignment(horizontal="center")
    ws.auto_filter.ref = f"A4:M{last}"
    for mr in list(ws.merged_cells.ranges):
        if mr.min_row == 2 and mr.max_row == 2:
            ws.unmerge_cells(str(mr)); ws.merge_cells("A2:M2")
    if "modelo de avión y la matrícula provienen" not in str(ws["A2"].value):
        ws["A2"] = (str(ws["A2"].value) + " El modelo de avión y la matrícula provienen de la API "
                    "del sitio de BER (api.flights.json); asientos y estela son valores típicos por modelo.")

    out = path if path.endswith("_con_aviones.xlsx") else re.sub(r"\.xlsx$", "_con_aviones.xlsx", path)
    wb.save(out)
    print(f"\nListo: {filled} vuelos con modelo, {pending} sin modelo. Guardado en {out}")
    if sin:
        print("Modelos sin clasificar (completalos en la hoja Aeronaves):", ", ".join(sin))
    if pending:
        print("Si faltan vuelos de una fecha futura, volvé a correr el script sobre el archivo "
              "_con_aviones.xlsx más cerca del día: completa solo lo que falte.")

if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else "vuelos_BER_2026-10-22_y_29.xlsx",
         sys.argv[2] if len(sys.argv) > 2 else None)